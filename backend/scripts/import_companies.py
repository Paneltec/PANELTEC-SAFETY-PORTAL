"""v160.3.9.19 — Companies XLSX ingestion (dedup by company_id)."""
from __future__ import annotations
import argparse, asyncio, hashlib, json, sys, uuid
from datetime import datetime, timezone
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))
from dotenv import load_dotenv
load_dotenv(BACKEND_ROOT / ".env")

import openpyxl  # noqa: E402
from db import db  # noqa: E402

DEFAULT_PATH = BACKEND_ROOT / "scripts" / "data" / "companies_source.xlsx"

COLUMN_MAP = {
    "A": "company_id", "B": "company", "C": "company_category",
    "D": "phone", "E": "suburb", "F": "state",
    "G": "general_email", "H": "archived_raw", "I": "account_type",
    "J": "company_classification",
}


def _now(): return datetime.now(timezone.utc).isoformat()


def _clean(v) -> str:
    if v is None: return ""
    if isinstance(v, (int, float)):
        s = str(v); return s[:-2] if s.endswith(".0") else s
    return str(v).strip()


def _yn(v) -> bool:
    return _clean(v).lower() in {"yes", "y", "true", "1"}


def _detect_header_row(ws) -> int:
    # Anchor on "Company category" so the section title "Companies" in row 1
    # doesn't false-positive.
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30),
                             values_only=False):
        for c in row:
            if _clean(c.value).lower() == "company category":
                return c.row
    return 7


def _content_hash(doc: dict) -> str:
    subset = {k: doc.get(k) for k in (
        "company_id", "company", "company_category", "phone", "suburb",
        "state", "general_email", "archived", "account_type",
        "company_classification",
    )}
    return hashlib.sha256(
        json.dumps(subset, sort_keys=True, ensure_ascii=False, default=str)
        .encode()).hexdigest()


def parse_workbook(path: Path) -> tuple[list[dict], set[str]]:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    header_row = _detect_header_row(ws)
    rows: list[dict] = []
    populated: set[str] = set()
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row,
                             values_only=False):
        cell = {c.column_letter: c for c in row}
        cid = _clean(cell.get("A").value) if cell.get("A") else ""
        if not cid:
            continue
        doc: dict = {"id": str(uuid.uuid4()), "deleted_at": None,
                     "company_id": cid, "archived": False}
        for letter, field in COLUMN_MAP.items():
            if field == "company_id":
                populated.add("company_id"); continue
            c = cell.get(letter)
            if not c or c.value in (None, ""):
                continue
            if field == "archived_raw":
                doc["archived"] = _yn(c.value)
                populated.add("archived")
                continue
            s = _clean(c.value)
            if s:
                doc[field] = s
                populated.add(field)
        rows.append(doc)
    return rows, populated


async def _resolve_actor_id(explicit):
    if explicit: return explicit
    u = await db.users.find_one({"email": "stephen@paneltec.com.au"},
                                 {"_id": 0, "id": 1})
    return u["id"] if u else "system"


async def ensure_indexes():
    await db.companies.create_index("company_id", unique=True)
    await db.companies.create_index("state")
    await db.companies.create_index("account_type")
    await db.companies_audit.create_index("company_id")


from _ingest_utils import merge_safe_upsert, detect_schema_shrink  # noqa: E402


async def upsert_rows(rows, actor_id, dry_run=False):
    inserted = updated = unchanged = 0
    for row in rows:
        row["content_hash"] = _content_hash(row)
        existing = await db.companies.find_one({"company_id": row["company_id"]}, {"_id": 0})
        if dry_run:
            if existing and existing.get("content_hash") == row["content_hash"]:
                unchanged += 1
            elif existing:
                updated += 1
            else:
                inserted += 1
            continue
        # Merge-safe: incoming empties never wipe existing fields (v160.3.9.21b)
        outcome = await merge_safe_upsert(
            collection=db.companies, audit_collection=db.companies_audit,
            key_field="company_id",
            incoming=row, existing=existing,
            actor_id=actor_id, content_hash=row["content_hash"],
        )
        if outcome == "inserted": inserted += 1
        elif outcome == "updated": updated += 1
        else: unchanged += 1
    return {"inserted": inserted, "updated": updated, "unchanged": unchanged}


async def main(args):
    path = Path(args.path).resolve()
    if not path.exists():
        print(f"[import] missing: {path}", file=sys.stderr); return 2
    print(f"[import] source={path} ({path.stat().st_size} bytes)")
    await ensure_indexes()
    rows, populated = parse_workbook(path)
    print(f"[import] parsed {len(rows)} rows; {len(populated)} populated columns")
    if not rows: return 3
    actor = await _resolve_actor_id(args.actor_id)
    dropped = await detect_schema_shrink(
        collection=db.companies, audit_collection=db.companies_audit,
        incoming_rows=rows, actor_id=actor)
    if dropped:
        print(f"[import] ⚠ schema shrink — {len(dropped)} column(s) missing: {dropped}")
        print("[import] existing values will be PRESERVED (merge-safe upsert).")
    stats = await upsert_rows(rows, actor, dry_run=args.dry_run)
    print(f"[import] inserted={stats['inserted']} updated={stats['updated']} unchanged={stats['unchanged']}")
    total = await db.companies.count_documents({"deleted_at": None})
    print(f"[import] live rows: {total}")
    print(f"[import] populated fields: {sorted(populated)}")
    return 0


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--path", default=str(DEFAULT_PATH))
    p.add_argument("--actor-id")
    p.add_argument("--dry-run", action="store_true")
    sys.exit(asyncio.run(main(p.parse_args())))
