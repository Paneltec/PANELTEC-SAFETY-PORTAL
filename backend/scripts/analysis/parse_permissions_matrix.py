"""v160.3.9.26 Phase-1 — Users & Permissions matrix parser (discovery-only).

Reads the Lucidity `User Matrix` export at
    /app/backend/scripts/inputs/users_permissions_matrix.xlsx

and emits two artefacts under /app/memory/permissions_redesign/:

  01_raw_taxonomy.json   — machine-readable dump of the parsed structure
  02_taxonomy_report.md  — human-readable summary

**Layout empirically confirmed** (see docstring output of the run):
  Row 6 = module group headings (12 merged ranges): `competency`, `induction`,
          `hr`, `incident`, `risk`, `inform`, `lucidityintranet`,
          `permittowork`, `contractor`, `asset`, `onsite`, `access`.
  Row 7 = sub-headings. Cols A..G are user metadata (Username, ID, First
          Name, Last Name, Email, Created At, Last Logged In). All later
          columns are per-module role sub-headings ("Administrator",
          "Column Configuration", "Manager", "Report Emailing", "General
          User", "Contractor Representative", "Read Only", "Responsible
          Manager", "Approver", "Manager (Creator)", "Mechanic Role",
          "Training / Inductions Only" and one long "Contractor
          Representative - only submit required documents").
  Rows 8..end (last row 119) = one user per row. Every non-metadata cell
          is either the string "Yes" or blank. A grant is "Yes"; blank is
          "no grant".

NO writes to Mongo, no backend/frontend edits. Discovery only.
"""

from __future__ import annotations

import json
import os
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

INPUT = Path("/app/backend/scripts/inputs/users_permissions_matrix.xlsx")
OUT_DIR = Path("/app/memory/permissions_redesign")
OUT_DIR.mkdir(parents=True, exist_ok=True)

METADATA_COL_HEADERS = {
    "Username", "ID", "First Name", "Last Name", "Email",
    "Created At", "Last Logged In",
}


def parse() -> dict:
    wb = openpyxl.load_workbook(INPUT, data_only=False)
    ws = wb.active

    # 1. Walk merged ranges on row 6 to build column -> module mapping.
    column_module: dict[int, str] = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 6 == rng.max_row:
            module = ws.cell(rng.min_row, rng.min_col).value
            if module is None:
                continue
            for c in range(rng.min_col, rng.max_col + 1):
                column_module[c] = str(module).strip()

    # 2. Row 7 sub-headings.
    column_role: dict[int, str] = {}
    metadata_columns: dict[int, str] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(7, c).value
        if v is None:
            continue
        v = str(v).strip()
        if v in METADATA_COL_HEADERS and c not in column_module:
            metadata_columns[c] = v
        else:
            column_role[c] = v

    # 3. Build the tidy list of (module, role, column) triples.
    permission_columns = []
    for c in sorted(column_role):
        if c in column_module:
            permission_columns.append({
                "column_index": c,
                "column_letter": get_column_letter(c),
                "module": column_module[c],
                "role": column_role[c],
            })

    # 4. Distinct roles per module (order of first appearance).
    modules: dict[str, list[str]] = {}
    for pc in permission_columns:
        modules.setdefault(pc["module"], [])
        if pc["role"] not in modules[pc["module"]]:
            modules[pc["module"]].append(pc["role"])

    # 5. Iterate user rows (row 8 through max_row where col A is populated).
    users = []
    for r in range(8, ws.max_row + 1):
        username = ws.cell(r, 1).value
        if username is None:
            continue
        user = {
            "row": r,
            "username": str(username).strip(),
            "id": ws.cell(r, 2).value,
            "first_name": ws.cell(r, 3).value,
            "last_name": ws.cell(r, 4).value,
            "email": ws.cell(r, 5).value,
            "created_at": _to_str(ws.cell(r, 6).value),
            "last_logged_in": _to_str(ws.cell(r, 7).value),
            "grants": [],
        }
        for pc in permission_columns:
            v = ws.cell(r, pc["column_index"]).value
            if v is None:
                continue
            token = str(v).strip()
            if not token:
                continue
            user["grants"].append({
                "module": pc["module"],
                "role": pc["role"],
                "value": token,
            })
        users.append(user)

    # 6. Distinct data-cell tokens per (module, role) column — should all be
    #    "Yes" or blank, but this proves it empirically.
    tokens_per_column: dict[tuple[str, str], Counter] = {}
    for pc in permission_columns:
        key = (pc["module"], pc["role"])
        c = Counter()
        for r in range(8, ws.max_row + 1):
            v = ws.cell(r, pc["column_index"]).value
            c[repr(v)] += 1
        tokens_per_column[key] = c

    return {
        "source_file": str(INPUT),
        "sheet": ws.title,
        "dimensions": {"rows": ws.max_row, "cols": ws.max_column},
        "merged_ranges_row6": [
            {
                "range": str(r),
                "module": ws.cell(r.min_row, r.min_col).value,
            }
            for r in sorted(ws.merged_cells.ranges, key=lambda x: x.min_col)
            if r.min_row == 6 == r.max_row and ws.cell(r.min_row, r.min_col).value
        ],
        "metadata_columns": [
            {"column_index": c, "column_letter": get_column_letter(c), "header": h}
            for c, h in sorted(metadata_columns.items())
        ],
        "permission_columns": permission_columns,
        "modules": modules,
        "users": users,
        "tokens_per_column": {
            f"{m}.{r}": dict(cnt) for (m, r), cnt in tokens_per_column.items()
        },
    }


def _to_str(v):
    if v is None:
        return None
    return str(v)


def write_json(data: dict) -> Path:
    p = OUT_DIR / "01_raw_taxonomy.json"
    p.write_text(json.dumps(data, indent=2, ensure_ascii=False))
    return p


def write_report(data: dict) -> Path:
    lines: list[str] = []
    lines.append("# 01 — Raw taxonomy (parsed from users_permissions_matrix.xlsx)")
    lines.append("")
    lines.append(
        f"Source: `{data['source_file']}` · sheet `{data['sheet']}` · "
        f"{data['dimensions']['rows']} rows × {data['dimensions']['cols']} cols"
    )
    lines.append("")
    lines.append(
        f"**Total users:** {len(data['users'])} · "
        f"**Modules:** {len(data['modules'])} · "
        f"**Permission columns:** {len(data['permission_columns'])}"
    )
    lines.append("")

    lines.append("## Row-6 module headings (merged)")
    lines.append("")
    lines.append("| Range | Module |")
    lines.append("|---|---|")
    for m in data["merged_ranges_row6"]:
        lines.append(f"| `{m['range']}` | `{m['module']}` |")
    lines.append("")

    lines.append("## Metadata columns (A–G)")
    lines.append("")
    lines.append("| Col | Header |")
    lines.append("|---|---|")
    for m in data["metadata_columns"]:
        lines.append(f"| `{m['column_letter']}` | {m['header']} |")
    lines.append("")

    lines.append("## Roles per module (order of first appearance)")
    lines.append("")
    for module, roles in data["modules"].items():
        lines.append(f"### `{module}`")
        for role in roles:
            lines.append(f"- {role}")
        lines.append("")

    lines.append("## Distinct cell-value tokens per (module, role) column")
    lines.append("")
    lines.append("Should be `'Yes'` or `None` throughout. Anything else = data-entry drift.")
    lines.append("")
    lines.append("| Column | Yes | Blank | Other |")
    lines.append("|---|---:|---:|---|")
    for key, counts in data["tokens_per_column"].items():
        yes = counts.get("'Yes'", 0)
        blank = counts.get("None", 0)
        other = {k: v for k, v in counts.items() if k not in ("'Yes'", "None")}
        other_str = ", ".join(f"{k}: {v}" for k, v in other.items()) if other else "—"
        lines.append(f"| `{key}` | {yes} | {blank} | {other_str} |")
    lines.append("")

    lines.append("## Users (username → count of grants)")
    lines.append("")
    lines.append("| Row | Username | Email | Grants |")
    lines.append("|---:|---|---|---:|")
    for u in data["users"]:
        lines.append(
            f"| {u['row']} | `{u['username']}` | "
            f"{u['email'] or ''} | {len(u['grants'])} |"
        )
    lines.append("")

    lines.append("## Full grant matrix (one row per user × module)")
    lines.append("")
    for u in data["users"]:
        if not u["grants"]:
            continue
        lines.append(f"### {u['username']} — {u['email'] or ''}")
        by_mod: dict[str, list[str]] = {}
        for g in u["grants"]:
            by_mod.setdefault(g["module"], []).append(g["role"])
        for mod, roles in by_mod.items():
            lines.append(f"- **{mod}**: {', '.join(roles)}")
        lines.append("")

    p = OUT_DIR / "02_taxonomy_report.md"
    p.write_text("\n".join(lines))
    return p


def main():
    data = parse()
    json_path = write_json(data)
    md_path = write_report(data)
    print(f"Wrote {json_path}  ({os.path.getsize(json_path)} bytes)")
    print(f"Wrote {md_path}   ({os.path.getsize(md_path)} bytes)")
    print("Users:", len(data["users"]))
    print("Modules:", len(data["modules"]))
    print("Permission columns:", len(data["permission_columns"]))


if __name__ == "__main__":
    main()
