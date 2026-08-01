"""v160.3.9.20 — Plant Maintenance XLSX ingestion.

Joins on `Registration No.` (col F) against `assets.rego_serial` with
case-insensitive uppercase-strip on both sides. Rows with a blank rego
are skipped; rows whose rego doesn't match any asset are imported with
`plant_id: null` so they surface in the "Unmatched" audit view.
"""
from __future__ import annotations
import argparse, asyncio, hashlib, json, sys, uuid
from datetime import datetime, timezone, date
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))
from dotenv import load_dotenv
load_dotenv(BACKEND_ROOT / ".env")

import openpyxl  # noqa: E402
from db import db  # noqa: E402

DEFAULT_PATH = BACKEND_ROOT / "scripts" / "data" / "plant_maintenance_source.xlsx"
SOURCE_URL = (
    "https://customer-assets-wrfwihn1.emergentagent.net/job_safety-pic-analyzer/"
    "artifacts/mvig4hzg_6a6859a73a4da-spreadsheet.xlsx"
)


def _ensure_source_downloaded(path: Path) -> None:
    """Download the canonical source XLSX iff not present on disk."""
    if path.exists() and path.stat().st_size > 0:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    import urllib.request
    print(f"[import] fetching source → {path}")
    req = urllib.request.Request(SOURCE_URL, headers={"User-Agent": "paneltec-import/1.0"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        path.write_bytes(resp.read())
    print(f"[import] downloaded {path.stat().st_size} bytes")

COLUMN_MAP = {
    "A": "maintenance_id", "B": "type", "C": "sub_type", "D": "description",
    "E": "manufacturer", "F": "registration_no", "G": "asset_code",
    "H": "maintenance_type", "I": "latest_usage_reading",
    "J": "due_at", "K": "due_date", "L": "maintenance_status",
    "M": "cost", "N": "company", "O": "date_completed",
    "P": "performed_by", "Q": "notes",
}


def _now(): return datetime.now(timezone.utc).isoformat()


def _clean(v) -> str:
    if v is None: return ""
    if isinstance(v, (int, float)):
        s = str(v); return s[:-2] if s.endswith(".0") else s
    return str(v).strip()


def _norm_rego(v) -> str:
    return _clean(v).upper().replace(" ", "")


def _parse_date(v):
    if v is None or (isinstance(v, str) and not v.strip()): return None
    if isinstance(v, datetime): return v.isoformat()
    if isinstance(v, date): return v.isoformat()
    s = str(v).strip()
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            dt = datetime.strptime(s, fmt)
            return dt.date().isoformat() if fmt in ("%d/%m/%Y", "%Y-%m-%d") else dt.isoformat()
        except ValueError:
            continue
    return s


def _detect_header_row(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=25, values_only=False):
        if _clean(row[0].value).lower() == "id" and any(
                _clean(c.value).lower() == "registration no." for c in row):
            return row[0].row
    return 7


def _content_hash(doc: dict) -> str:
    subset = {k: doc.get(k) for k in (
        "maintenance_id", "date_completed", "performed_by",
        "registration_matched", "maintenance_type",
    )}
    return hashlib.sha256(
        json.dumps(subset, sort_keys=True, ensure_ascii=False, default=str)
        .encode()).hexdigest()


async def _load_rego_index() -> dict:
    idx: dict = {}
    async for a in db.assets.find(
            {"deleted_at": None, "rego_serial": {"$ne": None, "$ne": ""}},
            {"_id": 0, "id": 1, "rego_serial": 1, "name": 1}):
        idx[_norm_rego(a["rego_serial"])] = a["id"]
    return idx


def parse_workbook(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    hdr = _detect_header_row(ws)
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=hdr + 1, max_row=ws.max_row,
                             values_only=False):
        cell = {c.column_letter: c for c in row}
        mid = _clean(cell.get("A").value) if cell.get("A") else ""
        if not mid: continue
        rego = _norm_rego(cell.get("F").value) if cell.get("F") else ""
        # Per user rule: skip only when rego is blank.
        if not rego:
            continue
        doc: dict = {
            "id": str(uuid.uuid4()),
            "registration_matched": rego,
            "plant_id": None,  # filled in during upsert if rego matches.
            "deleted_at": None,
        }
        for letter, field in COLUMN_MAP.items():
            c = cell.get(letter)
            if not c or c.value in (None, ""):
                continue
            if field == "registration_no":
                doc[field] = _clean(c.value)  # keep original casing/spacing
                continue
            if field in ("due_at", "due_date", "date_completed"):
                iso = _parse_date(c.value)
                if iso: doc[field] = iso
                continue
            s = _clean(c.value)
            if s: doc[field] = s
        rows.append(doc)
    return rows


async def _resolve_actor_id(explicit):
    if explicit: return explicit
    u = await db.users.find_one({"email": "stephen@paneltec.com.au"}, {"_id": 0, "id": 1})
    return u["id"] if u else "system"


async def ensure_indexes():
    await db.plant_maintenance.create_index("maintenance_id", unique=True)
    await db.plant_maintenance.create_index("plant_id")
    await db.plant_maintenance.create_index("registration_matched")
    await db.plant_maintenance.create_index("date_completed")
    await db.plant_maintenance_audit.create_index("maintenance_id")


async def upsert_rows(rows, actor_id, rego_idx, dry_run=False):
    inserted = updated = unchanged = matched = unmatched = 0
    now = _now()
    for row in rows:
        row["plant_id"] = rego_idx.get(row["registration_matched"])
        if row["plant_id"]: matched += 1
        else: unmatched += 1
        row["content_hash"] = _content_hash(row)
        existing = await db.plant_maintenance.find_one(
            {"maintenance_id": row["maintenance_id"]}, {"_id": 0})
        if existing and existing.get("content_hash") == row["content_hash"]:
            unchanged += 1; continue
        if existing:
            row["id"] = existing["id"]
            row["created_at"] = existing.get("created_at", now)
            row["imported_by"] = actor_id; row["imported_at"] = now
            row["updated_at"] = now
            if not dry_run:
                await db.plant_maintenance.replace_one(
                    {"maintenance_id": row["maintenance_id"]}, row, upsert=True)
                await db.plant_maintenance_audit.insert_one({
                    "id": str(uuid.uuid4()),
                    "maintenance_id": row["maintenance_id"],
                    "action": "update", "at": now, "actor_id": actor_id,
                    "before_hash": existing.get("content_hash"),
                    "after_hash": row["content_hash"],
                })
            updated += 1
        else:
            row["created_at"] = now
            row["imported_by"] = actor_id; row["imported_at"] = now
            if not dry_run:
                await db.plant_maintenance.insert_one(row)
                await db.plant_maintenance_audit.insert_one({
                    "id": str(uuid.uuid4()),
                    "maintenance_id": row["maintenance_id"],
                    "action": "insert", "at": now, "actor_id": actor_id,
                    "after_hash": row["content_hash"],
                })
            inserted += 1
    return {"inserted": inserted, "updated": updated, "unchanged": unchanged,
            "matched": matched, "unmatched": unmatched}


async def main(args):
    path = Path(args.path).resolve()
    _ensure_source_downloaded(path)
    if not path.exists(): print(f"[import] missing: {path}"); return 2
    print(f"[import] source={path} ({path.stat().st_size} bytes)")
    await ensure_indexes()
    rego_idx = await _load_rego_index()
    print(f"[import] {len(rego_idx)} assets indexed by rego")
    rows = parse_workbook(path)
    print(f"[import] {len(rows)} rows after skipping blank-rego rows")
    actor = await _resolve_actor_id(args.actor_id)
    stats = await upsert_rows(rows, actor, rego_idx, dry_run=args.dry_run)
    print(f"[import] inserted={stats['inserted']} updated={stats['updated']} "
          f"unchanged={stats['unchanged']}")
    print(f"[import] matched={stats['matched']} unmatched={stats['unmatched']}")
    return 0


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--path", default=str(DEFAULT_PATH))
    p.add_argument("--actor-id")
    p.add_argument("--dry-run", action="store_true")
    sys.exit(asyncio.run(main(p.parse_args())))
