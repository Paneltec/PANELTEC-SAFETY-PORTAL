"""v58.13.132lc — Smoke test: docx / xlsx / PDF through the new
in-process pipelines."""
import sys, io
sys.path.insert(0, '/app/backend')
from docx import Document
from openpyxl import Workbook
import fitz

results = {"docx_to_pdf": None, "xlsx_to_pdf": None,
           "pdf_text_extract": None, "pdf_ocr_lazy_degradation": None}

# Build a real .docx.
doc = Document()
doc.add_heading('Paneltec Civil — .132lc test doc', 0)
doc.add_paragraph('This is a body paragraph with some text. ' * 5)
tbl = doc.add_table(rows=3, cols=3)
for i, row in enumerate(tbl.rows):
    for j, cell in enumerate(row.cells):
        cell.text = f"R{i+1}C{j+1}"
buf = io.BytesIO(); doc.save(buf); docx_bytes = buf.getvalue()

# Build a real .xlsx.
wb = Workbook()
ws = wb.active; ws.title = "Compliance"
ws.append(["ID", "Title", "Status", "Owner"])
for i in range(1, 12):
    ws.append([f"CTR-{i:03d}", f"Contract {i}", "Active", "Stephen"])
ws2 = wb.create_sheet("Actions")
ws2.append(["Action", "Due"])
ws2.append(["Renew ISO 45001", "2026-11-01"])
buf = io.BytesIO(); wb.save(buf); xlsx_bytes = buf.getvalue()

# Build a real PDF with pymupdf.
pdf_doc = fitz.open()
page = pdf_doc.new_page()
page.insert_text((72, 100), "This is a text-layer PDF for testing ocr_pdf_to_text().", fontsize=12)
page2 = pdf_doc.new_page()
page2.insert_text((72, 100), "Page 2 body content. Compliance score: 98/100.", fontsize=11)
pdf_bytes = pdf_doc.tobytes(); pdf_doc.close()

# 1. docx → pdf.
from file_pdf import _docx_to_pdf, _xlsx_to_pdf, ocr_pdf_to_text
pdf_out, label = _docx_to_pdf(docx_bytes, "test.docx")
assert pdf_out[:5] == b"%PDF-", "docx output not valid PDF header"
assert len(pdf_out) > 500, f"docx PDF suspiciously small: {len(pdf_out)}"
results["docx_to_pdf"] = {"label": label, "bytes": len(pdf_out), "valid": True}

# 2. xlsx → pdf.
pdf_out = _xlsx_to_pdf(xlsx_bytes, "test.xlsx")
assert pdf_out[:5] == b"%PDF-", "xlsx output not valid PDF header"
assert len(pdf_out) > 800
results["xlsx_to_pdf"] = {"bytes": len(pdf_out), "valid": True}
# also verify it has both sheet titles.
with fitz.open(stream=pdf_out, filetype="pdf") as verify:
    all_text = " ".join(p.get_text() for p in verify)
assert "Compliance" in all_text, "missing Compliance sheet title"
assert "Actions" in all_text, "missing Actions sheet title"
results["xlsx_to_pdf"]["sheet_titles_found"] = True

# 3. PDF text extract (text-layer PDF, tesseract not needed).
tmp = "/tmp/_test.pdf"
with open(tmp, "wb") as f: f.write(pdf_bytes)
text = ocr_pdf_to_text(tmp)
assert "text-layer PDF" in text, f"text-layer extraction failed. got: {text!r}"
assert "Compliance score" in text
results["pdf_text_extract"] = {"chars": len(text), "found_content": True}

# 4. Verify lazy tesseract degradation on scanned PDF (no text layer).
scanned = fitz.open()
sp = scanned.new_page()
# Draw as bare rectangle without text — no text-layer.
sp.draw_rect(fitz.Rect(50,50,200,200), color=(0,0,0), fill=(1,1,1))
scanned_bytes = scanned.tobytes(); scanned.close()
with open("/tmp/_scanned.pdf", "wb") as f: f.write(scanned_bytes)
text = ocr_pdf_to_text("/tmp/_scanned.pdf")
# Since tesseract missing on this pod, should return "" — not crash.
results["pdf_ocr_lazy_degradation"] = {"returned_empty_no_crash": text == ""}

# 5. text_extraction with tesseract missing.
from text_extraction import extract_text
r = extract_text(scanned_bytes, mime="application/pdf", filename="scanned.pdf")
results["text_extraction_pdf_scanned"] = {"status": r["status"], "engine": r["engine"]}

import json
print(json.dumps(results, indent=2))
print("\nALL CHECKS PASSED" if all(
    r.get("valid") or r.get("found_content") or r.get("returned_empty_no_crash") or r.get("status")
    for r in results.values()) else "FAILURES")
