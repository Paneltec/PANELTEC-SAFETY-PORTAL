"""v160.3.9.15 — Incident Root Causes XLSX ingestion.

Mirrors `scripts/import_list_forms.py`. Reads
`/app/backend/scripts/data/incident_root_causes_source.xlsx`, parses each
data row into an `incident_root_causes` document, upserts by
`question_id`.

Run:
    cd /app/backend && python -m scripts.import_incident_root_causes
"""
from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from dotenv import load_dotenv
load_dotenv(BACKEND_ROOT / ".env")

import openpyxl  # noqa: E402

from db import db  # noqa: E402

DEFAULT_PATH = BACKEND_ROOT / "scripts" / "data" / "incident_root_causes_source.xlsx"


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        s = str(v)
        return s[:-2] if s.endswith(".0") else s
    return str(v).strip()


def _bool(v) -> bool:
    """Coerce Yes/True/1/checkbox-tick to True; anything else → False."""
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v != 0
    s = str(v).strip().lower()
    return s in {"yes", "y", "true", "t", "1", "✓", "✔", "checked"}


def _detect_header_row(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30),
                             values_only=False):
        v = _clean(row[0].value).lower().replace(" ", "")
        if v == "questionid":
            return row[0].row
    return 1


def _content_hash(doc: dict) -> str:
    subset = {k: doc.get(k) for k in (
        "question_id", "description", "contributing_factor",
        "parent_question_id", "has_action",
    )}
    return hashlib.sha256(
        json.dumps(subset, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()


def parse_workbook(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    header_row = _detect_header_row(ws)
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row,
                             values_only=False):
        cells = {c.column_letter: c for c in row}
        qid = _clean(cells.get("A").value) if cells.get("A") else ""
        if not qid:
            continue
        parent = _clean(cells.get("D").value) if cells.get("D") else ""
        doc = {
            "id": str(uuid.uuid4()),
            "question_id": qid,
            "description": _clean(cells.get("B").value) if cells.get("B") else "",
            "contributing_factor": _clean(cells.get("C").value) if cells.get("C") else "",
            "parent_question_id": parent or None,
            "has_action": _bool(cells.get("E").value) if cells.get("E") else False,
            "deleted_at": None,
        }
        rows.append(doc)
    return rows


async def _resolve_actor_id(explicit: str | None) -> str:
    if explicit:
        return explicit
    u = await db.users.find_one({"email": "stephen@paneltec.com.au"},
                                 {"_id": 0, "id": 1})
    return u["id"] if u else "system"


async def ensure_indexes() -> None:
    await db.incident_root_causes.create_index("question_id", unique=True)
    await db.incident_root_causes.create_index("contributing_factor")
    await db.incident_root_causes.create_index("parent_question_id")
    await db.incident_root_causes_audit.create_index("question_id")
    await db.incident_root_causes_audit.create_index("at")


async def upsert_rows(rows: list[dict], actor_id: str,
                      dry_run: bool = False) -> dict:
    inserted, updated, unchanged = 0, 0, 0
    now = _now_iso()
    for row in rows:
        row["content_hash"] = _content_hash(row)
        existing = await db.incident_root_causes.find_one(
            {"question_id": row["question_id"]}, {"_id": 0})
        if existing and existing.get("content_hash") == row["content_hash"]:
            unchanged += 1
            continue
        if existing:
            row["id"] = existing["id"]
            row["created_at"] = existing.get("created_at", now)
            row["imported_by"] = actor_id
            row["imported_at"] = now
            row["updated_at"] = now
            if not dry_run:
                await db.incident_root_causes.replace_one(
                    {"question_id": row["question_id"]}, row, upsert=True)
                await db.incident_root_causes_audit.insert_one({
                    "id": str(uuid.uuid4()),
                    "question_id": row["question_id"],
                    "action": "update", "at": now, "actor_id": actor_id,
                    "before_hash": existing.get("content_hash"),
                    "after_hash": row["content_hash"],
                })
            updated += 1
        else:
            row["created_at"] = now
            row["imported_by"] = actor_id
            row["imported_at"] = now
            if not dry_run:
                await db.incident_root_causes.insert_one(row)
                await db.incident_root_causes_audit.insert_one({
                    "id": str(uuid.uuid4()),
                    "question_id": row["question_id"],
                    "action": "insert", "at": now, "actor_id": actor_id,
                    "after_hash": row["content_hash"],
                })
            inserted += 1
    return {"inserted": inserted, "updated": updated, "unchanged": unchanged}


async def main(args) -> int:
    path = Path(args.path).resolve()
    if not path.exists():
        print(f"[import] source file missing: {path}", file=sys.stderr)
        return 2
    print(f"[import] source={path} ({path.stat().st_size} bytes)")
    await ensure_indexes()
    rows = parse_workbook(path)
    print(f"[import] parsed {len(rows)} rows")
    if not rows:
        return 3

    from collections import Counter
    cf_dist = Counter(r["contributing_factor"] or "<blank>" for r in rows)
    has_action_yes = sum(1 for r in rows if r["has_action"])
    parented = sum(1 for r in rows if r["parent_question_id"])
    print(f"[import] distinct contributing factors: {len(cf_dist)}")
    print(f"[import] has_action=Yes: {has_action_yes}/{len(rows)}")
    print(f"[import] parented rows: {parented}/{len(rows)}")

    actor_id = await _resolve_actor_id(args.actor_id)
    stats = await upsert_rows(rows, actor_id, dry_run=args.dry_run)
    print(f"[import] result → inserted={stats['inserted']}, "
          f"updated={stats['updated']}, unchanged={stats['unchanged']}")
    total = await db.incident_root_causes.count_documents({"deleted_at": None})
    print(f"[import] live rows after run: {total}")
    return 0


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--path", default=str(DEFAULT_PATH))
    p.add_argument("--actor-id")
    p.add_argument("--dry-run", action="store_true")
    sys.exit(asyncio.run(main(p.parse_args())))
