"""v160.3.9.17 — List Roles XLSX ingestion.

Source workbook stores `Capabilities` / `People` as per-role INTEGER
counts, not lists. We persist both the count fields AND reserved empty
`capabilities` / `people` string arrays so a future richer export can
populate them without a schema migration.
"""
from __future__ import annotations
import argparse, asyncio, hashlib, json, re, sys, uuid
from datetime import datetime, timezone
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))
from dotenv import load_dotenv
load_dotenv(BACKEND_ROOT / ".env")

import openpyxl  # noqa: E402
from db import db  # noqa: E402

DEFAULT_PATH = BACKEND_ROOT / "scripts" / "data" / "list_roles_source.xlsx"


def _now(): return datetime.now(timezone.utc).isoformat()


def _clean(v) -> str:
    if v is None: return ""
    if isinstance(v, (int, float)):
        s = str(v)
        return s[:-2] if s.endswith(".0") else s
    return str(v).strip()


def _int_or_zero(v) -> int:
    if v is None or v == "": return 0
    if isinstance(v, (int, float)): return int(v)
    try:
        return int(str(v).strip())
    except (ValueError, TypeError):
        return 0


def _split_list(v) -> list[str]:
    """Parse comma/semicolon/newline separated string → deduped list."""
    if v is None or (isinstance(v, (int, float))): return []
    s = str(v).strip()
    if not s: return []
    parts = re.split(r"[,;\n]", s)
    seen, out = set(), []
    for p in parts:
        p = p.strip()
        if p and p not in seen:
            seen.add(p); out.append(p)
    return out


def _detect_header_row(ws) -> int:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30),
                             values_only=False):
        for c in row:
            if _clean(c.value).lower() == "role title":
                return c.row
    return 1


def _content_hash(doc: dict) -> str:
    subset = {k: doc.get(k) for k in (
        "role_id", "role_title", "description",
        "capabilities_count", "people_count",
        "capabilities", "people",
    )}
    return hashlib.sha256(
        json.dumps(subset, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()


def parse_workbook(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb.active
    header_row = _detect_header_row(ws)
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row,
                             values_only=False):
        cell = {c.column_letter: c for c in row}
        role_id = _clean(cell.get("A").value) if cell.get("A") else ""
        if not role_id:
            continue
        cap_raw = cell.get("D").value if cell.get("D") else None
        ppl_raw = cell.get("E").value if cell.get("E") else None
        doc = {
            "id": str(uuid.uuid4()),
            "role_id": role_id,
            "role_title": _clean(cell.get("B").value) if cell.get("B") else "",
            "description": _clean(cell.get("C").value) if cell.get("C") else "",
            # v160.3.9.17 — source stores counts; reserved string arrays
            # default to [] so richer future exports can populate them
            # without a schema migration.
            "capabilities_count": _int_or_zero(cap_raw),
            "people_count": _int_or_zero(ppl_raw),
            "capabilities": _split_list(cap_raw),
            "people": _split_list(ppl_raw),
            "deleted_at": None,
        }
        rows.append(doc)
    return rows


async def _resolve_actor_id(explicit):
    if explicit: return explicit
    u = await db.users.find_one({"email": "stephen@paneltec.com.au"},
                                 {"_id": 0, "id": 1})
    return u["id"] if u else "system"


async def ensure_indexes() -> None:
    await db.list_roles.create_index("role_id", unique=True)
    await db.list_roles.create_index("role_title")
    await db.list_roles_audit.create_index("role_id")


from _ingest_utils import merge_safe_upsert, detect_schema_shrink  # noqa: E402


async def upsert_rows(rows, actor_id, dry_run=False):
    inserted = updated = unchanged = 0
    for row in rows:
        row["content_hash"] = _content_hash(row)
        existing = await db.list_roles.find_one(
            {"role_id": row["role_id"]}, {"_id": 0})
        if dry_run:
            if existing and existing.get("content_hash") == row["content_hash"]:
                unchanged += 1
            elif existing:
                updated += 1
            else:
                inserted += 1
            continue
        # Merge-safe: incoming empties never wipe existing fields (v160.3.9.21b)
        outcome = await merge_safe_upsert(
            collection=db.list_roles, audit_collection=db.list_roles_audit,
            key_field="role_id",
            incoming=row, existing=existing,
            actor_id=actor_id, content_hash=row["content_hash"],
        )
        if outcome == "inserted": inserted += 1
        elif outcome == "updated": updated += 1
        else: unchanged += 1
    return {"inserted": inserted, "updated": updated, "unchanged": unchanged}


async def main(args) -> int:
    path = Path(args.path).resolve()
    if not path.exists():
        print(f"[import] missing: {path}", file=sys.stderr); return 2
    print(f"[import] source={path} ({path.stat().st_size} bytes)")
    await ensure_indexes()
    rows = parse_workbook(path)
    print(f"[import] parsed {len(rows)} rows")
    if not rows: return 3
    tot_cap = sum(r["capabilities_count"] for r in rows)
    tot_ppl = sum(r["people_count"] for r in rows)
    print(f"[import] sum capabilities_count: {tot_cap}, sum people_count: {tot_ppl}")
    actor = await _resolve_actor_id(args.actor_id)
    dropped = await detect_schema_shrink(
        collection=db.list_roles, audit_collection=db.list_roles_audit,
        incoming_rows=rows, actor_id=actor)
    if dropped:
        print(f"[import] ⚠ schema shrink — {len(dropped)} column(s) missing: {dropped}")
        print("[import] existing values will be PRESERVED (merge-safe upsert).")
    stats = await upsert_rows(rows, actor, dry_run=args.dry_run)
    print(f"[import] inserted={stats['inserted']} updated={stats['updated']} unchanged={stats['unchanged']}")
    total = await db.list_roles.count_documents({"deleted_at": None})
    print(f"[import] live rows: {total}")
    return 0


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--path", default=str(DEFAULT_PATH))
    p.add_argument("--actor-id")
    p.add_argument("--dry-run", action="store_true")
    sys.exit(asyncio.run(main(p.parse_args())))
