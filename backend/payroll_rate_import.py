"""Preview legacy pay rates against current Simpro identities. Never creates workers."""
import io,json,math,re,zipfile
from collections import defaultdict
from fastapi import APIRouter,Depends,HTTPException,UploadFile,File,Form
from permissions import require_permission
from db import db

router=APIRouter(prefix='/import',tags=['payroll-import'])

async def workbook(file):
    from openpyxl import load_workbook
    raw=await file.read(5_000_001)
    if len(raw)>5_000_000:raise HTTPException(422,'Workbook must be smaller than 5 MB')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>25_000_000:raise ValueError('Workbook is too large')
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        if any(s.max_row>10000 or s.max_column>150 for s in book):raise ValueError('Workbook dimensions exceed the import limit')
        result={s.title:list(s.values) for s in book};book.close();return result
    except HTTPException:raise
    except Exception:raise HTTPException(422,'Could not read this XLSX workbook; check its format and size')

def number(value):
    try:
        n=float(value or 0)
        if not math.isfinite(n) or n<0 or n>10000000:raise ValueError()
        return n
    except (TypeError,ValueError):raise ValueError('Invalid numeric pay rate')

def norm(value):return ' '.join(str(value or '').casefold().split())

def rate_profile(source,categories,existing,pay_basis=None):
    from payroll_workbench import Profile
    profile=Profile(**existing).model_dump();rates={};warnings=[]
    for key,value in source.items():
        match=re.match(r'^PC(\d+)_',str(key))
        if match:rates[match[1]]=number(value)
    byname={c['PayCategoryName']:c for c in categories.values()}
    def get(name):return rates.get(str(byname.get(name,{}).get('Id')),0)
    salary=get('Salary');base=get('Casual Ordinary Hours');ordinary=get('Permanent Ordinary Hours')
    if pay_basis not in (None,'annual_salary','hourly'):raise ValueError('Invalid pay basis choice')
    if salary and (base or ordinary):
        if pay_basis=='annual_salary':
            base=ordinary=0
            warnings.append('Annual salary selected; the source hourly rate is excluded from the primary pay rate.')
        elif pay_basis=='hourly':salary=0
        else:raise ValueError('Both salary and hourly rates are populated; choose the intended pay basis manually')
    if salary:
        if byname['Salary'].get('RateUnit')!='Annually':raise ValueError('Salary must be expressed annually')
        profile.update(pay_basis='annual_salary',annual_salary=salary,casual_rates=None)
    elif norm(source.get('Pay Rate Template'))=='casual' or base:
        if not base:raise ValueError('Casual ordinary base rate is missing')
        cat=byname['Casual Ordinary Hours']
        if cat.get('RateUnit')!='Hourly':raise ValueError('Casual ordinary rate must be hourly')
        loading=(number(cat.get('RateLoadingMultiplier'))-1)*100
        if loading<0 or loading>100:raise ValueError('Casual loading is outside supported limits')
        profile.update(employment_type='casual',pay_basis='hourly',annual_salary=0,
            casual_rates={'base_rate':base,'loading_percent':loading,
                'ot1':get('Casual - Overtime 1.5') or None,
                'ot2':get('Casual - Overtime 2.0') or None,
                'holiday_work':get('Casual - Public Holiday Worked') or None,'night':None})
        warnings.append('Casual leave accrual remains governed by employee employment type; the source category flag does not enable paid annual or sick leave.')
    elif ordinary:
        if byname['Permanent Ordinary Hours'].get('RateUnit')!='Hourly':raise ValueError('Ordinary rate must be hourly')
        profile.update(pay_basis='hourly',annual_salary=0,hourly_rate=ordinary,casual_rates=None)
        if profile['employment_type']=='casual':raise ValueError('Permanent source rate conflicts with saved casual employment type')
    else:raise ValueError('No positive ordinary rate or annual salary found')
    applied={str(byname[n]['Id']) for n in ('Salary','Casual Ordinary Hours','Permanent Ordinary Hours','Casual - Overtime 1.5','Casual - Overtime 2.0','Casual - Public Holiday Worked') if n in byname}
    other=[{'category_id':k,'name':categories.get(k,{}).get('PayCategoryName',k),'rate':v,'unit':categories.get(k,{}).get('RateUnit','Unknown')} for k,v in rates.items() if v and k not in applied]
    profile['conditions_reviewed']=False
    return Profile(**profile).model_dump(),warnings,other

@router.post('/preview')
async def preview(rates:UploadFile=File(...),categories:UploadFile=File(...),aliases:str=Form('{}'),pay_bases:str=Form('{}'),user=Depends(require_permission('payroll','edit'))):
    from payroll_roster import roster
    from payroll_employee_records import public
    try:
        aliases=json.loads(aliases)
        if not isinstance(aliases,dict) or len(aliases)>500 or any(not isinstance(k,str) or not isinstance(v,str) for k,v in aliases.items()):raise ValueError()
    except Exception:raise HTTPException(422,'Employee matches must map source names to current worker IDs')
    try:
        pay_bases=json.loads(pay_bases)
        if not isinstance(pay_bases,dict) or len(pay_bases)>500 or any(not isinstance(k,str) or v not in ('annual_salary','hourly') for k,v in pay_bases.items()):raise ValueError()
    except Exception:raise HTTPException(422,'Pay basis choices must map current worker IDs to salary or hourly')
    books=await workbook(rates);cats=await workbook(categories)
    rows=books.get('Export',[]);category_rows=cats.get('Export',[])
    if not rows or not {'Employee Id','First Name','Surname'}.issubset(rows[0]):raise HTTPException(422,'Select the Employee Pay Rates extract')
    if not category_rows or not {'Id','PayCategoryName','RateUnit','RateLoadingMultiplier'}.issubset(category_rows[0]):raise HTTPException(422,'Select the Pay Categories extract')
    definitions={str(r[0]):dict(zip(category_rows[0],r)) for r in category_rows[1:] if r[0] is not None}
    current=await roster(user['org_id']);names=defaultdict(list);ids={w['id']:w for w in current}
    for w in current:names[norm(w['name'])].append(w['id'])
    ready=[];unmatched=[];seen=set()
    for index,values in enumerate(rows[1:],2):
        source=dict(zip(rows[0],values));name=f"{source.get('First Name') or ''} {source.get('Surname') or ''}".strip()
        if not name:continue
        matches=[aliases[name]] if name in aliases else names[norm(name)]
        if len(matches)!=1 or matches[0] not in ids:
            unmatched.append({'source_name':name,'source_row':index,'reason':'No unique current Simpro employee match'});continue
        wid=matches[0]
        if wid in seen:raise HTTPException(422,'Multiple source rows map to the same employee; resolve duplicates before importing')
        seen.add(wid)
        record=public(await db.pay_employee_records.find_one({'_id':f"{user['org_id']}:{wid}"}))
        try:profile,warnings,other=rate_profile(source,definitions,record['profile'],pay_bases.get(wid))
        except ValueError as e:
            unmatched.append({'source_name':name,'source_row':index,'worker_id':wid,'reason':str(e)});continue
        ready.append({'worker_id':wid,'name':ids[wid]['name'],'source_name':name,'revision':record['revision'],'profile':profile,'warnings':warnings,'other_rates':other,'previous_profile':record['profile']})
    return {'rows':ready,'unmatched':unmatched,'workers':[{'id':w['id'],'name':w['name']} for w in current],
        'unchanged':[w['name'] for w in current if w['id'] not in {r['worker_id'] for r in ready}],
        'scope':'Preview only. Ordinary rates, annual salaries, casual loading and explicit casual overtime/public-holiday rates can be applied. Other category rates remain listed for separate setup. Saved tax, fund, bank and opening balances are preserved.'}
